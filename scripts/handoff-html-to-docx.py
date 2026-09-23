#!/usr/bin/env python3
"""
Convert a handoff HTML document into a Word .docx with real Word tables.

Why this exists: `textutil -convert docx` (macOS) produces valid output but
flattens every table into paragraphs, which loses the comparison tables the
handoff plan depends on. This script preserves them.

Requires: python-docx, beautifulsoup4
    python3 -m venv .venv && .venv/bin/pip install python-docx beautifulsoup4

Usage:
    .venv/bin/python scripts/handoff-html-to-docx.py \
        handoff/gerotech-handoff-plan.html handoff/gerotech-handoff-plan.docx
"""

import sys
from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

ORANGE = RGBColor(0xB4, 0x50, 0x00)   # --clr-orange-deep, readable on white
INK = RGBColor(0x0D, 0x0D, 0x0D)


def add_runs(paragraph, node, bold=False, italic=False, link=False):
    """Walk inline children, preserving <strong>/<em>/<code>/<a>."""
    for child in node.children:
        if isinstance(child, NavigableString):
            text = str(child)
            if not text.strip() and not paragraph.runs:
                text = text.lstrip()
            if text:
                run = paragraph.add_run(text)
                run.bold = bold
                run.italic = italic
                if link:
                    run.font.color.rgb = ORANGE
                    run.underline = True
        elif isinstance(child, Tag):
            name = child.name.lower()
            if name == 'br':
                paragraph.add_run().add_break()
            elif name in ('strong', 'b'):
                add_runs(paragraph, child, True, italic, link)
            elif name in ('em', 'i'):
                add_runs(paragraph, child, bold, True, link)
            elif name == 'a':
                add_runs(paragraph, child, bold, italic, True)
                href = child.get('href')
                if href and href.startswith('http'):
                    tail = paragraph.add_run(f" ({href})")
                    tail.font.size = Pt(9)
                    tail.font.color.rgb = RGBColor(0x6B, 0x6B, 0x76)
            elif name == 'code':
                run = paragraph.add_run(child.get_text())
                run.font.name = 'Menlo'
                run.font.size = Pt(9.5)
            else:
                add_runs(paragraph, child, bold, italic, link)


def add_table(doc, table_tag):
    rows = table_tag.find_all('tr')
    if not rows:
        return
    ncols = max(len(r.find_all(['td', 'th'])) for r in rows)
    tbl = doc.add_table(rows=0, cols=ncols)
    tbl.style = 'Table Grid'
    for r in rows:
        cells = r.find_all(['td', 'th'])
        if not cells:
            continue
        row = tbl.add_row()
        is_header = cells[0].name == 'th'
        for i, cell in enumerate(cells[:ncols]):
            para = row.cells[i].paragraphs[0]
            add_runs(para, cell, bold=is_header)
            for run in para.runs:
                run.font.size = Pt(9.5)
                if is_header:
                    run.font.color.rgb = INK
                elif i == 0:
                    run.bold = True
    doc.add_paragraph()


def convert(src, dst):
    soup = BeautifulSoup(open(src, encoding='utf-8').read(), 'html.parser')
    body = soup.body or soup
    doc = Document()

    # Base document styling
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(8)

    for el in body.find_all(recursive=False):
        name = el.name.lower() if el.name else None
        if name is None:
            continue

        if name == 'h1':
            p = doc.add_paragraph()
            run = p.add_run(el.get_text().strip())
            run.bold = True
            run.font.size = Pt(20)
            run.font.color.rgb = INK
        elif name in ('h2', 'h3', 'h4'):
            size = {'h2': 14, 'h3': 12, 'h4': 11}[name]
            p = doc.add_paragraph()
            run = p.add_run(el.get_text().strip())
            run.bold = True
            run.font.size = Pt(size)
            run.font.color.rgb = ORANGE if name == 'h2' else INK
            p.paragraph_format.space_before = Pt(14)
        elif name == 'p':
            p = doc.add_paragraph()
            add_runs(p, el)
            if el.get('class') and 'meta' in el.get('class'):
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(0x6B, 0x6B, 0x76)
        elif name in ('ul', 'ol'):
            for li in el.find_all('li', recursive=False):
                p = doc.add_paragraph(style='List Number' if name == 'ol' else 'List Bullet')
                add_runs(p, li)
        elif name == 'table':
            add_table(doc, el)
        elif name == 'hr':
            doc.add_paragraph()

    doc.save(dst)
    print(f"wrote {dst}")


if __name__ == '__main__':
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    convert(sys.argv[1], sys.argv[2])
